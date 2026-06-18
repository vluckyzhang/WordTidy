from __future__ import annotations

import re
from pathlib import Path

from docx import Document


MD_HEADING_RE = re.compile(r"^(#{1,8})\s+(.+?)\s*$")


def text_to_docx(input_path: Path, output_path: Path, source_type: str) -> Path:
    text = _read_text(input_path)
    document = Document()

    if source_type == ".md":
        _append_markdown(document, text)
    else:
        _append_plain_text(document, text)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _append_plain_text(document: Document, text: str) -> None:
    for block in _paragraph_blocks(text):
        document.add_paragraph(block)


def _append_markdown(document: Document, text: str) -> None:
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        document.add_paragraph(_strip_basic_markdown(" ".join(paragraph_lines)))
        paragraph_lines.clear()

    for raw_line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw_line.strip()
        if not line:
            flush_paragraph()
            continue

        heading_match = MD_HEADING_RE.match(line)
        if heading_match:
            flush_paragraph()
            level = min(8, len(heading_match.group(1)))
            paragraph = document.add_paragraph(heading_match.group(2).strip())
            paragraph.style = f"Heading {level}"
            continue

        if line.startswith(("- ", "* ", "+ ")):
            flush_paragraph()
            document.add_paragraph(_strip_basic_markdown(line[2:].strip()), style="List Bullet")
            continue

        ordered_match = re.match(r"^\d+[.)]\s+(.+)$", line)
        if ordered_match:
            flush_paragraph()
            document.add_paragraph(_strip_basic_markdown(ordered_match.group(1).strip()), style="List Number")
            continue

        paragraph_lines.append(line)

    flush_paragraph()


def _paragraph_blocks(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    blocks = [block.strip() for block in re.split(r"\n\s*\n", normalized)]
    return [block for block in blocks if block]


def _strip_basic_markdown(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"__([^_]+)__", r"\1", text)
    text = re.sub(r"_([^_]+)_", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.strip()
