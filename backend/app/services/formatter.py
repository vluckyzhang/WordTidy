from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.schemas.rules import (
    DirectoryLevelRule,
    FontRule,
    FormattingRules,
    FormulaRule,
    ImageRule,
    ParagraphRule,
    TableRule,
)


ParagraphLabels = dict[int, str]

CAPTION_RE = re.compile(r"^\s*(图|表|Figure|Table)\s*[\d一二三四五六七八九十IVXivx\-_.]*")
FORMULA_DELIMITER_RE = re.compile(r"^\s*(?:\$\$(?P<dollar>.+?)\$\$|\\\[(?P<bracket>.+?)\\\]|\\\((?P<paren>.+?)\\\))\s*$", re.S)
FORMULA_HINT_RE = re.compile(r"(\\frac|\\sum|\\int|\\sqrt|\\left|\\right|[_^{}]|[A-Za-z0-9]\s*=\s*[-+*/\\A-Za-z0-9(){}\[\]\s]+)")
HEADING3_RE = re.compile(r"^\s*(\d+[.．]\d+[.．]\d+|[（(][a-zA-Z0-9一二三四五六七八九十]+[）)])")
HEADING2_RE = re.compile(r"^\s*(\d+[.．]\d+|（[一二三四五六七八九十\d]+）)")
HEADING1_RE = re.compile(
    r"^\s*(第[一二三四五六七八九十百千万\d]+[章节]|[一二三四五六七八九十]+[、．.]|摘要$|Abstract$|引言$|结论$|参考文献$)",
    re.IGNORECASE,
)


def format_docx(
    input_path: Path,
    output_path: Path,
    rules: FormattingRules,
    paragraph_labels: ParagraphLabels | None = None,
    insert_directories: bool = False,
) -> list[str]:
    warnings: list[str] = []
    document = Document(str(input_path))

    if rules.update_fields_on_open:
        _enable_update_fields_on_open(document)

    _apply_page_setup(document, rules)
    _move_captions_above_targets(document)
    _format_styles(document, rules)
    _format_paragraphs(document, rules, paragraph_labels or {})
    _format_tables(document, rules.table)
    _format_existing_toc_styles(document, rules)

    if insert_directories:
        _insert_directory_fields(document, rules)
        warnings.append("已插入目录/图目录/表目录字段；首次打开 Word 时请允许更新域以生成页码。")

    warnings.extend(_configure_page_numbers(document, rules))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return warnings


def _apply_page_setup(document: Document, rules: FormattingRules) -> None:
    for section in document.sections:
        if rules.page.size == "A4":
            section.page_width = Cm(21)
            section.page_height = Cm(29.7)
        section.top_margin = Cm(rules.page.top_cm)
        section.bottom_margin = Cm(rules.page.bottom_cm)
        section.left_margin = Cm(rules.page.left_cm)
        section.right_margin = Cm(rules.page.right_cm)
        section.header_distance = Cm(rules.page.header_cm)
        section.footer_distance = Cm(rules.page.footer_cm)


def _format_styles(document: Document, rules: FormattingRules) -> None:
    style_rules = {
        "Normal": rules.body,
        "Body Text": rules.body,
        "Caption": rules.caption,
    }
    for level, rule in _iter_heading_rules(rules):
        style_rules[f"Heading {level}"] = rule
    for style_name, rule in style_rules.items():
        if style_name not in document.styles:
            continue
        style = document.styles[style_name]
        _apply_font_to_style(style, rule)
        if hasattr(style, "paragraph_format"):
            _apply_paragraph_format(style.paragraph_format, rule)


def _format_paragraphs(document: Document, rules: FormattingRules, labels: ParagraphLabels) -> None:
    for index, paragraph in enumerate(document.paragraphs):
        if not paragraph.text.strip() and not _paragraph_has_image(paragraph):
            continue

        if _paragraph_has_image(paragraph):
            _apply_image_paragraph_format(paragraph, rules.image)
            continue

        if rules.formula.enabled and _paragraph_is_formula(paragraph, rules.formula):
            _apply_formula_paragraph_format(paragraph, rules.formula)
            continue

        label = labels.get(index) or _classify_paragraph(paragraph)
        rule = _rule_for_label(label, rules)
        _apply_paragraph_format(paragraph.paragraph_format, rule)
        _apply_font_to_paragraph(paragraph, rule)


def _format_tables(document: Document, rule: TableRule) -> None:
    for table in document.tables:
        _clear_table_formatting(table)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        _set_table_width_percent(table, rule.width_percent)
        _apply_three_line_borders(table, rule)
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.alignment = _alignment(rule.cell_alignment)
                    _apply_line_spacing_format(paragraph.paragraph_format, rule.line_spacing_type, rule.line_spacing_value)
                    _apply_font_to_paragraph(paragraph, rule)


def _format_existing_toc_styles(document: Document, rules: FormattingRules) -> None:
    directory_rules: dict[str, DirectoryLevelRule] = {
        "TOC 1": rules.directory.level1,
        "TOC 2": rules.directory.level2,
        "TOC 3": rules.directory.level3,
    }
    for style_name, rule in directory_rules.items():
        if style_name not in document.styles:
            continue
        style = document.styles[style_name]
        _apply_font_to_style(style, rule)
        style.paragraph_format.left_indent = _chars_to_pt(rule.indent_chars, rule.size_pt)
        style.paragraph_format.alignment = _alignment(rule.alignment)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(24)

    if "TOC Heading" in document.styles:
        _apply_font_to_style(document.styles["TOC Heading"], rules.directory.title)
        document.styles["TOC Heading"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _insert_directory_fields(document: Document, rules: FormattingRules) -> None:
    first = document.paragraphs[0] if document.paragraphs else document.add_paragraph()
    leading_text = "\n".join(p.text.strip() for p in document.paragraphs[:12])
    if "目录" in leading_text and "图目录" in leading_text and "表目录" in leading_text:
        return

    max_level = max(rules.directory.enabled_levels or [3])
    _insert_directory_block(first, "目录", rf'TOC \o "1-{max_level}" \h \z \u', rules)
    _insert_directory_block(first, "图目录", r'TOC \h \z \c "图"', rules)
    _insert_directory_block(first, "表目录", r'TOC \h \z \c "表"', rules)
    page_break = first.insert_paragraph_before()
    page_break.add_run().add_break(WD_BREAK.PAGE)


def _insert_directory_block(anchor: Paragraph, title: str, instruction: str, rules: FormattingRules) -> None:
    title_paragraph = anchor.insert_paragraph_before(title)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _apply_font_to_paragraph(title_paragraph, rules.directory.title)

    field_paragraph = anchor.insert_paragraph_before()
    field_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _add_field(field_paragraph, instruction)

    spacer = anchor.insert_paragraph_before()
    spacer.paragraph_format.space_after = Pt(0)


def _configure_page_numbers(document: Document, rules: FormattingRules) -> list[str]:
    warnings: list[str] = []
    sections = list(document.sections)
    if not sections:
        return warnings

    if len(sections) == 1:
        _set_footer_page_number(sections[0], rules, "decimal")
        warnings.append("当前文档只有一个节，已设置正文阿拉伯页码；目录罗马页码需要模板或手动分节后再自动应用。")
        return warnings

    for index, section in enumerate(sections):
        style = "roman" if index == 0 else "decimal"
        _set_footer_page_number(section, rules, style)
        _set_section_page_number_format(section, style, start=1)
        section.start_type = WD_SECTION_START.NEW_PAGE
    return warnings


def _set_footer_page_number(section, rules: FormattingRules, style: str) -> None:
    paragraph = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    paragraph.clear()
    paragraph.alignment = _alignment(rules.page_number.alignment)
    _add_field(paragraph, "PAGE")
    for run in paragraph.runs:
        _apply_font_to_run(run, rules.page_number)
    _set_section_page_number_format(section, style, start=1)


def _set_section_page_number_format(section, style: str, start: int = 1) -> None:
    section_properties = section._sectPr
    page_number_type = section_properties.find(qn("w:pgNumType"))
    if page_number_type is None:
        page_number_type = OxmlElement("w:pgNumType")
        section_properties.append(page_number_type)
    page_number_type.set(qn("w:start"), str(start))
    page_number_type.set(qn("w:fmt"), "lowerRoman" if style == "roman" else "decimal")


def _move_captions_above_targets(document: Document) -> None:
    body = document.element.body
    children = list(body.iterchildren())
    for child in children:
        if child.tag != qn("w:p"):
            continue
        paragraph = Paragraph(child, document)
        if not CAPTION_RE.match(paragraph.text.strip()):
            continue
        previous = child.getprevious()
        if previous is None:
            continue
        if previous.tag == qn("w:tbl") or _paragraph_element_has_image(previous):
            previous.addprevious(child)


def _classify_paragraph(paragraph: Paragraph) -> str:
    text = paragraph.text.strip()
    style_name = paragraph.style.name.lower() if paragraph.style is not None else ""
    for level in range(1, 9):
        if f"heading {level}" in style_name or f"标题 {level}" in style_name:
            return f"heading{level}"
    if "caption" in style_name or "题注" in style_name or CAPTION_RE.match(text):
        return "caption"
    numbered_level = _numbered_heading_level(text)
    if numbered_level:
        return f"heading{numbered_level}"
    if HEADING3_RE.match(text):
        return "heading3"
    if HEADING2_RE.match(text):
        return "heading2"
    if HEADING1_RE.match(text):
        return "heading1"
    return "body"


def _rule_for_label(label: str, rules: FormattingRules) -> ParagraphRule:
    heading_match = re.fullmatch(r"heading([1-8])", label)
    if heading_match:
        return _heading_rule(rules, int(heading_match.group(1)))
    if label == "caption":
        return rules.caption
    return rules.body


def _apply_font_to_style(style, rule: FontRule) -> None:
    font = style.font
    font.name = rule.ascii
    font.size = Pt(rule.size_pt)
    if not rule.preserve_emphasis:
        font.bold = rule.bold
        font.italic = rule.italic
    style.element.rPr.rFonts.set(qn("w:eastAsia"), rule.east_asia)
    style.element.rPr.rFonts.set(qn("w:ascii"), rule.ascii)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), rule.ascii)


def _apply_font_to_paragraph(paragraph: Paragraph, rule: FontRule) -> None:
    runs = paragraph.runs
    if not runs and paragraph.text:
        runs = [paragraph.add_run()]
    for run in runs:
        _apply_font_to_run(run, rule)


def _apply_font_to_run(run, rule: FontRule) -> None:
    run.font.name = rule.ascii
    run.font.size = Pt(rule.size_pt)
    if not rule.preserve_emphasis:
        run.font.bold = rule.bold
        run.font.italic = rule.italic
    run_element = run._element
    run_properties = run_element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.append(fonts)
    fonts.set(qn("w:eastAsia"), rule.east_asia)
    fonts.set(qn("w:ascii"), rule.ascii)
    fonts.set(qn("w:hAnsi"), rule.ascii)


def _apply_paragraph_format(paragraph_format, rule: ParagraphRule) -> None:
    paragraph_format.space_before = Pt(rule.space_before_pt)
    paragraph_format.space_after = Pt(rule.space_after_pt)
    paragraph_format.alignment = _alignment(rule.alignment)
    paragraph_format.first_line_indent = _chars_to_pt(rule.first_line_indent_chars, rule.size_pt)
    _apply_line_spacing_format(paragraph_format, rule.line_spacing_type, rule.line_spacing_value)


def _apply_image_paragraph_format(paragraph: Paragraph, rule: ImageRule) -> None:
    paragraph.paragraph_format.space_before = Pt(rule.space_before_pt)
    paragraph.paragraph_format.space_after = Pt(rule.space_after_pt)
    paragraph.paragraph_format.alignment = _alignment(rule.alignment)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    _apply_line_spacing_format(paragraph.paragraph_format, rule.line_spacing_type, rule.line_spacing_value)


def _apply_formula_paragraph_format(paragraph: Paragraph, rule: FormulaRule) -> None:
    formula_text = _extract_formula_text(paragraph.text, rule)
    paragraph.clear()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.alignment = _alignment(rule.alignment)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    if rule.format == "professional":
        _add_office_math(paragraph, formula_text)
        return

    run = paragraph.add_run(formula_text)
    run.font.name = rule.font
    run.font.size = Pt(rule.size_pt)
    run_element = run._element
    run_properties = run_element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.append(fonts)
    fonts.set(qn("w:ascii"), rule.font)
    fonts.set(qn("w:hAnsi"), rule.font)
    fonts.set(qn("w:eastAsia"), rule.font)
    fonts.set(qn("w:cs"), rule.font)


def _apply_line_spacing_format(paragraph_format, spacing_type: str, value: float) -> None:
    if spacing_type == "fixed":
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph_format.line_spacing = Pt(value)
    elif spacing_type == "single":
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph_format.line_spacing = None
    elif spacing_type == "one_point_five":
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        paragraph_format.line_spacing = None
    elif spacing_type == "double":
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        paragraph_format.line_spacing = None
    else:
        paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        paragraph_format.line_spacing = value


def _iter_heading_rules(rules: FormattingRules):
    for level in range(1, 9):
        rule = getattr(rules.headings, f"level{level}", None)
        if rule is not None:
            yield level, rule


def _heading_rule(rules: FormattingRules, level: int) -> ParagraphRule:
    rule = getattr(rules.headings, f"level{level}", None)
    if rule is not None:
        return rule
    return rules.body


def _numbered_heading_level(text: str) -> int | None:
    match = re.match(r"^\s*(\d+(?:[.．]\d+){0,7})(?:\s|[、．.]|$)", text)
    if not match:
        return None
    marker = match.group(1)
    return min(8, marker.count(".") + marker.count("．") + 1)


def _paragraph_is_formula(paragraph: Paragraph, rule: FormulaRule) -> bool:
    text = paragraph.text.strip()
    if not text:
        return False
    if FORMULA_DELIMITER_RE.match(text):
        return True
    if not rule.ai_enhanced_detection:
        return False
    if len(text) > 220:
        return False
    return bool(FORMULA_HINT_RE.search(text))


def _extract_formula_text(text: str, rule: FormulaRule) -> str:
    stripped = text.strip()
    if rule.strip_delimiters:
        match = FORMULA_DELIMITER_RE.match(stripped)
        if match:
            return next(value for value in match.groupdict().values() if value is not None).strip()
    return stripped


def _add_office_math(paragraph: Paragraph, formula_text: str) -> None:
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = formula_text
    math_run.append(math_text)
    math.append(math_run)
    paragraph._p.append(math)


def _set_table_width_percent(table: Table, width_percent: int) -> None:
    table_properties = table._tbl.tblPr
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:type"), "pct")
    table_width.set(qn("w:w"), str(width_percent * 50))


def _apply_three_line_borders(table: Table, rule: TableRule) -> None:
    row_count = len(table.rows)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            _set_cell_borders(cell, top=("nil", 0), bottom=("nil", 0), left=("nil", 0), right=("nil", 0))
            if row_index == 0:
                _set_cell_borders(
                    cell,
                    top=("single", _border_width(rule.thick_border_pt)),
                    bottom=("single", _border_width(rule.thin_border_pt)),
                )
            if row_index == row_count - 1:
                _set_cell_borders(cell, bottom=("single", _border_width(rule.thick_border_pt)))


def _clear_table_formatting(table: Table) -> None:
    try:
        table.style = None
    except Exception:
        pass

    table_properties = table._tbl.tblPr
    if table_properties is not None:
        _remove_children(
            table_properties,
            [
                "w:tblStyle",
                "w:tblLook",
                "w:tblBorders",
                "w:tblCellMar",
                "w:tblW",
                "w:jc",
                "w:tblLayout",
                "w:shd",
            ],
        )

    for row in table.rows:
        for cell in row.cells:
            cell_properties = cell._tc.get_or_add_tcPr()
            _remove_children(
                cell_properties,
                [
                    "w:tcBorders",
                    "w:shd",
                    "w:tcMar",
                    "w:vAlign",
                    "w:textDirection",
                    "w:noWrap",
                ],
            )
            for paragraph in cell.paragraphs:
                _clear_paragraph_direct_formatting(paragraph)
                for run in paragraph.runs:
                    _clear_run_direct_formatting(run)


def _clear_paragraph_direct_formatting(paragraph: Paragraph) -> None:
    paragraph_properties = paragraph._p.pPr
    if paragraph_properties is None:
        return
    _remove_children(
        paragraph_properties,
        [
            "w:pStyle",
            "w:jc",
            "w:spacing",
            "w:ind",
            "w:contextualSpacing",
            "w:keepNext",
            "w:keepLines",
        ],
    )


def _clear_run_direct_formatting(run) -> None:
    run_properties = run._r.rPr
    if run_properties is not None:
        run._r.remove(run_properties)


def _remove_children(parent, tags: list[str]) -> None:
    for tag in tags:
        element = parent.find(qn(tag))
        while element is not None:
            parent.remove(element)
            element = parent.find(qn(tag))


def _set_cell_borders(cell, **borders: tuple[str, int]) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    tc_borders = cell_properties.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        cell_properties.append(tc_borders)
    for edge, (value, size) in borders.items():
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), value)
        if size:
            element.set(qn("w:sz"), str(size))
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "000000")


def _border_width(point: float) -> int:
    return max(2, int(round(point * 8)))


def _add_field(paragraph: Paragraph, instruction: str) -> None:
    run = paragraph.add_run()

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run._r.append(instr)

    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)

    placeholder = OxmlElement("w:t")
    placeholder.text = "请更新域"
    run._r.append(placeholder)

    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _enable_update_fields_on_open(document: Document) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def _paragraph_has_image(paragraph: Paragraph) -> bool:
    return _paragraph_element_has_image(paragraph._p)


def _paragraph_element_has_image(element) -> bool:
    return bool(element.xpath(".//w:drawing | .//w:pict"))


def _alignment(value: str):
    return {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(value, WD_ALIGN_PARAGRAPH.LEFT)


def _chars_to_pt(chars: float, font_size_pt: float):
    return Pt(chars * font_size_pt) if chars else Pt(0)
